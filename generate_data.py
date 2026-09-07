import openpyxl
import json
import re
import os

def run():
    excel_path = os.path.join('Data', 'Mapping.xlsx')
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb['Mapping with Reference to NRB']

    iso_standard = {
        'Clause 4.1': {'code': 'Clause 4.1', 'title': 'Understanding the organization and its context', 'category': 'Clauses', 'theme': 'Clause 4: Context of the Organization'},
        'Clause 4.2': {'code': 'Clause 4.2', 'title': 'Understanding the needs and expectations of interested parties', 'category': 'Clauses', 'theme': 'Clause 4: Context of the Organization'},
        'Clause 4.3': {'code': 'Clause 4.3', 'title': 'Determining the scope of the information security management system', 'category': 'Clauses', 'theme': 'Clause 4: Context of the Organization'},
        'Clause 4.4': {'code': 'Clause 4.4', 'title': 'Information security management system', 'category': 'Clauses', 'theme': 'Clause 4: Context of the Organization'},
        'Clause 5.1': {'code': 'Clause 5.1', 'title': 'Leadership and commitment', 'category': 'Clauses', 'theme': 'Clause 5: Leadership'},
        'Clause 5.2': {'code': 'Clause 5.2', 'title': 'Policy', 'category': 'Clauses', 'theme': 'Clause 5: Leadership'},
        'Clause 5.3': {'code': 'Clause 5.3', 'title': 'Organizational roles, responsibilities and authorities', 'category': 'Clauses', 'theme': 'Clause 5: Leadership'},
        'Clause 6.1.1': {'code': 'Clause 6.1.1', 'title': 'Actions to address risks and opportunities - General', 'category': 'Clauses', 'theme': 'Clause 6: Planning'},
        'Clause 6.1.2': {'code': 'Clause 6.1.2', 'title': 'Information security risk assessment', 'category': 'Clauses', 'theme': 'Clause 6: Planning'},
        'Clause 6.1.3': {'code': 'Clause 6.1.3', 'title': 'Information security risk treatment', 'category': 'Clauses', 'theme': 'Clause 6: Planning'},
        'Clause 6.2': {'code': 'Clause 6.2', 'title': 'Information security objectives and planning to achieve them', 'category': 'Clauses', 'theme': 'Clause 6: Planning'},
        'Clause 6.3': {'code': 'Clause 6.3', 'title': 'Planning of changes', 'category': 'Clauses', 'theme': 'Clause 6: Planning'},
        'Clause 7.1': {'code': 'Clause 7.1', 'title': 'Resources', 'category': 'Clauses', 'theme': 'Clause 7: Support'},
        'Clause 7.2': {'code': 'Clause 7.2', 'title': 'Competence', 'category': 'Clauses', 'theme': 'Clause 7: Support'},
        'Clause 7.3': {'code': 'Clause 7.3', 'title': 'Awareness', 'category': 'Clauses', 'theme': 'Clause 7: Support'},
        'Clause 7.4': {'code': 'Clause 7.4', 'title': 'Communication', 'category': 'Clauses', 'theme': 'Clause 7: Support'},
        'Clause 7.5.1': {'code': 'Clause 7.5.1', 'title': 'Documented information - General', 'category': 'Clauses', 'theme': 'Clause 7: Support'},
        'Clause 7.5.2': {'code': 'Clause 7.5.2', 'title': 'Creating and updating', 'category': 'Clauses', 'theme': 'Clause 7: Support'},
        'Clause 7.5.3': {'code': 'Clause 7.5.3', 'title': 'Control of documented information', 'category': 'Clauses', 'theme': 'Clause 7: Support'},
        'Clause 8.1': {'code': 'Clause 8.1', 'title': 'Operational planning and control', 'category': 'Clauses', 'theme': 'Clause 8: Operation'},
        'Clause 8.2': {'code': 'Clause 8.2', 'title': 'Information security risk assessment (Operation)', 'category': 'Clauses', 'theme': 'Clause 8: Operation'},
        'Clause 8.3': {'code': 'Clause 8.3', 'title': 'Information security risk treatment (Operation)', 'category': 'Clauses', 'theme': 'Clause 8: Operation'},
        'Clause 9.1': {'code': 'Clause 9.1', 'title': 'Monitoring, measurement, analysis and evaluation', 'category': 'Clauses', 'theme': 'Clause 9: Performance Evaluation'},
        'Clause 9.2.1': {'code': 'Clause 9.2.1', 'title': 'Internal audit - General', 'category': 'Clauses', 'theme': 'Clause 9: Performance Evaluation'},
        'Clause 9.2.2': {'code': 'Clause 9.2.2', 'title': 'Internal audit programme', 'category': 'Clauses', 'theme': 'Clause 9: Performance Evaluation'},
        'Clause 9.3.1': {'code': 'Clause 9.3.1', 'title': 'Management review - General', 'category': 'Clauses', 'theme': 'Clause 9: Performance Evaluation'},
        'Clause 9.3.2': {'code': 'Clause 9.3.2', 'title': 'Management review inputs', 'category': 'Clauses', 'theme': 'Clause 9: Performance Evaluation'},
        'Clause 9.3.3': {'code': 'Clause 9.3.3', 'title': 'Management review results', 'category': 'Clauses', 'theme': 'Clause 9: Performance Evaluation'},
        'Clause 10.1': {'code': 'Clause 10.1', 'title': 'Continual improvement', 'category': 'Clauses', 'theme': 'Clause 10: Improvement'},
        'Clause 10.2': {'code': 'Clause 10.2', 'title': 'Nonconformity and corrective action', 'category': 'Clauses', 'theme': 'Clause 10: Improvement'},

        'A.5.1': {'code': 'A.5.1', 'title': 'Policies for information security', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.2': {'code': 'A.5.2', 'title': 'Information security roles and responsibilities', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.3': {'code': 'A.5.3', 'title': 'Segregation of duties', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.4': {'code': 'A.5.4', 'title': 'Management responsibilities', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.5': {'code': 'A.5.5', 'title': 'Contact with authorities', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.6': {'code': 'A.5.6', 'title': 'Contact with special interest groups', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.7': {'code': 'A.5.7', 'title': 'Threat intelligence', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.8': {'code': 'A.5.8', 'title': 'Information security in project management', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.9': {'code': 'A.5.9', 'title': 'Inventory of information and other associated assets', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.10': {'code': 'A.5.10', 'title': 'Acceptable use of information and other associated assets', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.11': {'code': 'A.5.11', 'title': 'Return of assets', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.12': {'code': 'A.5.12', 'title': 'Classification of information', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.13': {'code': 'A.5.13', 'title': 'Labelling of information', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.14': {'code': 'A.5.14', 'title': 'Information transfer', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.15': {'code': 'A.5.15', 'title': 'Access control', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.16': {'code': 'A.5.16', 'title': 'Identity management', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.17': {'code': 'A.5.17', 'title': 'Authentication information', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.18': {'code': 'A.5.18', 'title': 'Access rights', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.19': {'code': 'A.5.19', 'title': 'Information security in supplier relationships', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.20': {'code': 'A.5.20', 'title': 'Addressing information security within supplier agreements', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.21': {'code': 'A.5.21', 'title': 'Managing information security in the ICT supply chain', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.22': {'code': 'A.5.22', 'title': 'Monitoring, review and change management of supplier services', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.23': {'code': 'A.5.23', 'title': 'Information security for use of cloud services', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.24': {'code': 'A.5.24', 'title': 'Information security incident management planning and preparation', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.25': {'code': 'A.5.25', 'title': 'Assessment and decision on information security events', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.26': {'code': 'A.5.26', 'title': 'Response to information security incidents', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.27': {'code': 'A.5.27', 'title': 'Learning from information security incidents', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.28': {'code': 'A.5.28', 'title': 'Collection of evidence', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.29': {'code': 'A.5.29', 'title': 'Information security during disruption', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.30': {'code': 'A.5.30', 'title': 'ICT readiness for business continuity', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.31': {'code': 'A.5.31', 'title': 'Legal, statutory, regulatory and contractual requirements', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.32': {'code': 'A.5.32', 'title': 'Intellectual property rights', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.33': {'code': 'A.5.33', 'title': 'Protection of records', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.34': {'code': 'A.5.34', 'title': 'Privacy and protection of PII', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.35': {'code': 'A.5.35', 'title': 'Independent review of information security', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.36': {'code': 'A.5.36', 'title': 'Compliance with policies and standards for information security', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},
        'A.5.37': {'code': 'A.5.37', 'title': 'Documented operating procedures', 'category': 'Annex A', 'theme': 'A.5 Organizational Controls'},

        'A.6.1': {'code': 'A.6.1', 'title': 'Screening', 'category': 'Annex A', 'theme': 'A.6 People Controls'},
        'A.6.2': {'code': 'A.6.2', 'title': 'Terms and conditions of employment', 'category': 'Annex A', 'theme': 'A.6 People Controls'},
        'A.6.3': {'code': 'A.6.3', 'title': 'Information security awareness, education and training', 'category': 'Annex A', 'theme': 'A.6 People Controls'},
        'A.6.4': {'code': 'A.6.4', 'title': 'Disciplinary process', 'category': 'Annex A', 'theme': 'A.6 People Controls'},
        'A.6.5': {'code': 'A.6.5', 'title': 'Responsibilities after termination or change of employment', 'category': 'Annex A', 'theme': 'A.6 People Controls'},
        'A.6.6': {'code': 'A.6.6', 'title': 'Confidentiality or non-disclosure agreements', 'category': 'Annex A', 'theme': 'A.6 People Controls'},
        'A.6.7': {'code': 'A.6.7', 'title': 'Remote working', 'category': 'Annex A', 'theme': 'A.6 People Controls'},
        'A.6.8': {'code': 'A.6.8', 'title': 'Information security event reporting', 'category': 'Annex A', 'theme': 'A.6 People Controls'},

        'A.7.1': {'code': 'A.7.1', 'title': 'Physical security perimeters', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.2': {'code': 'A.7.2', 'title': 'Physical entry', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.3': {'code': 'A.7.3', 'title': 'Securing offices, rooms and facilities', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.4': {'code': 'A.7.4', 'title': 'Physical security monitoring', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.5': {'code': 'A.7.5', 'title': 'Protecting against physical and environmental threats', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.6': {'code': 'A.7.6', 'title': 'Working in secure areas', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.7': {'code': 'A.7.7', 'title': 'Clear desk and clear screen', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.8': {'code': 'A.7.8', 'title': 'Equipment siting and protection', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.9': {'code': 'A.7.9', 'title': 'Security of assets off-premises', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.10': {'code': 'A.7.10', 'title': 'Storage media', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.11': {'code': 'A.7.11', 'title': 'Supporting utilities', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.12': {'code': 'A.7.12', 'title': 'Cabling security', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.13': {'code': 'A.7.13', 'title': 'Equipment maintenance', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},
        'A.7.14': {'code': 'A.7.14', 'title': 'Secure disposal or re-use of equipment', 'category': 'Annex A', 'theme': 'A.7 Physical Controls'},

        'A.8.1': {'code': 'A.8.1', 'title': 'User endpoint devices', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.2': {'code': 'A.8.2', 'title': 'Privileged access rights', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.3': {'code': 'A.8.3', 'title': 'Information access restriction', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.4': {'code': 'A.8.4', 'title': 'Access to source code', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.5': {'code': 'A.8.5', 'title': 'Secure authentication', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.6': {'code': 'A.8.6', 'title': 'Capacity management', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.7': {'code': 'A.8.7', 'title': 'Protection against malware', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.8': {'code': 'A.8.8', 'title': 'Management of technical vulnerabilities', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.9': {'code': 'A.8.9', 'title': 'Configuration management', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.10': {'code': 'A.8.10', 'title': 'Information deletion', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.11': {'code': 'A.8.11', 'title': 'Data masking', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.12': {'code': 'A.8.12', 'title': 'Data leakage prevention', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.13': {'code': 'A.8.13', 'title': 'Information backup', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.14': {'code': 'A.8.14', 'title': 'Redundancy of information processing facilities', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.15': {'code': 'A.8.15', 'title': 'Logging', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.16': {'code': 'A.8.16', 'title': 'Monitoring activities', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.17': {'code': 'A.8.17', 'title': 'Clock synchronization', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.18': {'code': 'A.8.18', 'title': 'Use of privileged utility programs', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.19': {'code': 'A.8.19', 'title': 'Installation of software on operational systems', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.20': {'code': 'A.8.20', 'title': 'Networks security', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.21': {'code': 'A.8.21', 'title': 'Security of network services', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.22': {'code': 'A.8.22', 'title': 'Segregation of networks', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.23': {'code': 'A.8.23', 'title': 'Web filtering', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.24': {'code': 'A.8.24', 'title': 'Use of cryptography', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.25': {'code': 'A.8.25', 'title': 'Secure development life cycle', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.26': {'code': 'A.8.26', 'title': 'Application security requirements', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.27': {'code': 'A.8.27', 'title': 'Secure system architecture and engineering principles', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.28': {'code': 'A.8.28', 'title': 'Secure coding', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.29': {'code': 'A.8.29', 'title': 'Security testing in development and acceptance', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.30': {'code': 'A.8.30', 'title': 'Outsourced development', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.31': {'code': 'A.8.31', 'title': 'Separation of development, test and production environments', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.32': {'code': 'A.8.32', 'title': 'Change management', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.33': {'code': 'A.8.33', 'title': 'Test information', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'},
        'A.8.34': {'code': 'A.8.34', 'title': 'Protection of information systems during audit testing', 'category': 'Annex A', 'theme': 'A.8 Technological Controls'}
    }

    raw_rows = []
    for r in range(1, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        raw_rows.append(vals)

    nrb_sections = [
        {'id': 1, 'name': 'IT Governance', 'desc': 'Board oversight, strategic alignment, IT organizational structure, and information security roles.'},
        {'id': 2, 'name': 'Information Security', 'desc': 'Security policies, risk assessment, access controls, hardening, encryption, and electronic channels.'},
        {'id': 3, 'name': 'Information Security Education', 'desc': 'Awareness programs for employees, vendors, and customer security education.'},
        {'id': 4, 'name': 'Information Disclosure and Grievance Handling', 'desc': 'Customer dispute handling SLAs, privacy notice publication, and fee transparency.'},
        {'id': 5, 'name': 'Outsourcing Management', 'desc': 'Vendor risk assessment, contracts, inspection rights, and source code escrow.'},
        {'id': 6, 'name': 'IT Operations', 'desc': 'Segregation of duties, formal change management, dual custody, and operational procedures.'},
        {'id': 7, 'name': 'Information Systems Acquisition, Development and Implementation', 'desc': 'SDLC security requirements, forensic audit trails, and source code review.'},
        {'id': 8, 'name': 'Business Continuity and Disaster Recovery Planning', 'desc': 'BCP policy, RPO/RTO parameters, DR drills, DC-DR replication, and high-availability architecture.'},
        {'id': 9, 'name': 'IS Audit', 'desc': 'Audit team resources, independent scope, and IS audit governance.'},
        {'id': 10, 'name': 'Fraud Management', 'desc': 'Cyber attack reporting to NRB and customer fraud education.'}
    ]

    nrb_items = []
    links = []
    
    current_sec_id = 0
    current_sec_name = ''
    sec_item_counters = {i: 0 for i in range(1, 11)}

    for row_idx, r in enumerate(raw_rows[5:103], start=6):
        col0 = str(r[0]).strip() if r[0] is not None else ''
        if col0.startswith('Section '):
            parts = col0.split('—')
            sec_num_str = parts[0].replace('Section', '').strip()
            current_sec_id = int(sec_num_str)
            current_sec_name = parts[1].strip() if len(parts) > 1 else col0
            continue
        
        if r[1] and r[4]:
            sec_item_counters[current_sec_id] += 1
            ref_str = f'{current_sec_id}.{sec_item_counters[current_sec_id]}'
            nrb_id = f'NRB-{ref_str}'
            
            nrb_section = str(r[1]).strip()
            nrb_req = str(r[2]).strip()
            iso_raw = str(r[3]).strip() if r[3] else ''
            overlap = str(r[4]).strip()
            notes = str(r[5]).strip() if r[5] else ''

            if overlap == 'None - NRB unique':
                status_type = 'none'
                status_label = 'NRB Unique (Not in ISO)'
                status_color = '#ef4444' # Red
            elif overlap == 'Partial':
                status_type = 'partial'
                status_label = 'Partial Overlap'
                status_color = '#f59e0b' # Yellow
            elif overlap == 'Full / Partial':
                status_type = 'mixed'
                status_label = 'Full / Partial'
                status_color = '#10b981' # Green primary
            else: # 'Full'
                status_type = 'full'
                status_label = 'Full Overlap'
                status_color = '#10b981' # Green

            item_obj = {
                'id': nrb_id,
                'ref': ref_str,
                'sectionId': current_sec_id,
                'sectionName': current_sec_name,
                'requirement': nrb_req,
                'isoRaw': iso_raw,
                'overlap': overlap,
                'statusType': status_type,
                'statusLabel': status_label,
                'statusColor': status_color,
                'notes': notes,
                'connectedIso': []
            }

            if status_type != 'none' and iso_raw and iso_raw != '—':
                lines = [l.strip() for l in iso_raw.splitlines() if l.strip()]
                idx_line = 0
                while idx_line < len(lines):
                    line = lines[idx_line]
                    if line.startswith('[') and line.endswith(']'):
                        idx_line += 1
                        continue
                    if line == 'to' and idx_line > 0 and idx_line + 1 < len(lines):
                        idx_line += 1
                        continue
                    if idx_line + 2 < len(lines) and lines[idx_line+1] == 'to':
                        start_code = lines[idx_line].split('-')[0].strip()
                        end_code = lines[idx_line+2].split('-')[0].strip()
                        if start_code == 'A.7.1' and end_code == 'A.7.5':
                            for k in range(1, 6):
                                iso_key = f'A.7.{k}'
                                link_overlap = 'Partial' if 'partial' in line.lower() or overlap == 'Partial' else 'Full'
                                item_obj['connectedIso'].append({
                                    'isoCode': iso_key,
                                    'overlap': link_overlap
                                })
                                links.append({
                                    'source': nrb_id,
                                    'target': f'ISO-{iso_key.replace(" ", "_")}',
                                    'isoCode': iso_key,
                                    'nrbRef': ref_str,
                                    'overlap': link_overlap,
                                    'color': '#10b981' if link_overlap == 'Full' else '#f59e0b'
                                })
                        idx_line += 3
                        continue

                    m = re.match(r'^(Clause\s+[0-9\.]+|A\.[0-9\.]+)\s*(?:-\s*([^(\n]+))?(?:\s*\(([^)]+)\))?', line)
                    if m:
                        iso_code = m.group(1).strip()
                        is_partial = False
                        if m.group(3) and 'partial' in m.group(3).lower():
                            is_partial = True
                        elif 'partial' in line.lower():
                            is_partial = True
                        elif overlap == 'Partial':
                            is_partial = True
                        
                        link_overlap = 'Partial' if is_partial else 'Full'
                        item_obj['connectedIso'].append({
                            'isoCode': iso_code,
                            'overlap': link_overlap
                        })
                        links.append({
                            'source': nrb_id,
                            'target': f'ISO-{iso_code.replace(" ", "_")}',
                            'isoCode': iso_code,
                            'nrbRef': ref_str,
                            'overlap': link_overlap,
                            'color': '#10b981' if link_overlap == 'Full' else '#f59e0b'
                        })
                    idx_line += 1

            nrb_items.append(item_obj)

    iso_items = []
    iso_connections = {}
    for l in links:
        code = l['isoCode']
        if code not in iso_connections:
            iso_connections[code] = []
        iso_connections[code].append({
            'nrbId': l['source'],
            'nrbRef': l['nrbRef'],
            'overlap': l['overlap'],
            'color': l['color']
        })

    for code, info in iso_standard.items():
        conns = iso_connections.get(code, [])
        is_mapped = len(conns) > 0
        has_full = any(c['overlap'] == 'Full' for c in conns)
        
        if not is_mapped:
            iso_status = 'unmapped'
            iso_color = '#64748b'
            iso_badge = 'Not Mapped in NRB'
        elif has_full:
            iso_status = 'full'
            iso_color = '#10b981'
            iso_badge = f'{len(conns)} NRB Link{"s" if len(conns)>1 else ""}'
        else:
            iso_status = 'partial'
            iso_color = '#f59e0b'
            iso_badge = f'{len(conns)} NRB Link{"s" if len(conns)>1 else ""} (Partial)'

        iso_items.append({
            'id': f'ISO-{code.replace(" ", "_")}',
            'code': code,
            'title': info['title'],
            'category': info['category'],
            'theme': info['theme'],
            'isMapped': is_mapped,
            'status': iso_status,
            'badge': iso_badge,
            'color': iso_color,
            'connectedNrb': conns
        })

    stats = {
        'totalNrb': len(nrb_items),
        'fullNrb': sum(1 for item in nrb_items if item['statusType'] in ['full', 'mixed']),
        'partialNrb': sum(1 for item in nrb_items if item['statusType'] == 'partial'),
        'noneNrb': sum(1 for item in nrb_items if item['statusType'] == 'none'),
        'totalIsoControls': len(iso_items),
        'mappedIsoControls': sum(1 for item in iso_items if item['isMapped']),
        'totalLinks': len(links)
    }

    out_data = {
        'stats': stats,
        'nrbSections': nrb_sections,
        'nrbItems': nrb_items,
        'isoItems': iso_items,
        'links': links
    }

    with open('data.js', 'w', encoding='utf-8') as f:
        f.write('const MAPPING_DATA = ' + json.dumps(out_data, indent=2, ensure_ascii=False) + ';\n')

    print('Successfully generated data.js!')
    print(json.dumps(stats, indent=2))

if __name__ == '__main__':
    run()
